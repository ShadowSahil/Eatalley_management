"""
Reading the Google Sheet expense tracker and filing its rows.

Each month tab ("August", "October", ...) has the columns
Date | Amount | By | To | Type | Comment(s) | Mode.

Where a row goes:
  * paid to one of our vendors (Vendors page: JS Foods, Sunrise/SFS, Naveen,
    Mokafarms, Hyperpure) -> skipped by the monthly import; their bills are
    added under Inventory. (The one-off Aug–Oct command files them as bills.)
  * Type "Bill" / "Rent", and staff room rent  -> Fixed costs
  * Type "Salary": helpers / labour / day-wage   -> "Adhoc helper" payments;
    main staff are skipped (their pay is recorded on the Staff page)
  * Type "Equipments" / "Interior"             -> Daily expenses (kitchen
    equipment, repairs), so they count in the P&L
  * everything else                            -> Daily expenses; Inventory rows
    get the item named in the comment (item list seeded by
    expenses/migrations/0004_sheet_items.py)
The one-off Aug–Oct command covers the set-up period instead
(opening_period=True): Equipments / Interior are Opening costs, and every
salary is filed (monthly staff on the Staff page, everyone else Adhoc helper).
"""

import hashlib
import io
import re
import urllib.error
import urllib.request
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from django.db import transaction

from core.models import Payer, SheetImport
from core.periods import add_months, month_end
from expenses.models import DailyExpense, ExpenseItem
from fixedcosts.models import FixedCost, FixedCostHead
from inventory.models import Purchase, Vendor
from opening.models import OpeningCost
from payroll.models import SalaryPayment, Staff

ZERO = Decimal(0)
BUSINESS_PAYER = "Eatalley"

ADHOC = ("Adhoc helper", "Daily / ad-hoc helpers and labour")

# --------------------------------------------------------------- daily expense items

# Phrases rewritten to single tokens before matching (order matters).
PHRASES = [
    (r"ice\s*cream", "icecream"),
    (r"milk\s*maid|milkmade|milkmaid|condensed?\s*milk|condesed\s*milk|condensed", "milkmaid"),
    (r"coconut\s*water", "coconutwater"),
    (r"cold\s*drinks?", "colddrink"),
    (r"diet\s*coke", "coke"),
    (r"ice\s*cubes?|\bcubes?\b", "ice"),
    (r"ice\s*tray", "tray"),
    (r"milk\s*straw", "straw"),
    (r"water\s*bottles?", "waterbottle"),
    (r"coffee\s*sachet", "coffeesachet"),
    (r"kitchen\s*towel", "kitchentowel"),
    (r"display\s*packet", "packet"),
    (r"cherry\s*tomato", "cherrytomato"),
    (r"peanut\s*butter", "peanutbutter"),
    (r"wrapping\s*paper", "wrapping"),
    (r"room\s*fresh(e)?ner", "freshener"),
]
# Everything bought under "Inventory" outside the vendors above. Keep the
# names in step with expenses/migrations/0004_sheet_items.py.
ITEMS = [  # name, category, unit, token regex
    ("Milk", "dairy", "L", r"milk"),
    ("Paneer", "dairy", "kg", r"paneer"),
    ("Condensed milk (Milkmaid)", "dairy", "pcs", r"milkmaid"),
    ("Dahi", "dairy", "kg", r"dahi|curd"),
    ("Adhoc cheese", "dairy", "pkt", r"cheese|parm|parmesan|cheddar"),  # top-ups outside JS Foods
    ("Ice", "beverages", "bag", r"ice"),
    ("Ice cream", "beverages", "pcs", r"icecream"),
    ("Cold Drink", "beverages", "btl", r"colddrink|coke|sprite|soda"),
    ("Coconut Water", "beverages", "pcs", r"coconutwater"),
    ("Coffee sachets", "beverages", "pkt", r"coffeesachet"),
    ("Protein powder", "beverages", "kg", r"protein"),
    ("Water bottles", "beverages", "btl", r"waterbottle"),
    ("Syrups", "beverages", "btl", r"syrup"),
    ("Watermelon", "veg_fruit", "kg", r"watermelon"),
    ("Banana", "veg_fruit", "dozen", r"banana"),
    ("Cherries", "veg_fruit", "pkt", r"cheery|cherry"),
    ("Lemon", "veg_fruit", "kg", r"lemon|nembu|nimbu"),
    ("Mint", "veg_fruit", "pkt", r"mint"),
    ("Parsley", "veg_fruit", "pkt", r"parsley"),
    ("Basil", "veg_fruit", "pkt", r"basil"),
    ("Broccoli", "veg_fruit", "kg", r"broccoli"),
    ("Mushroom", "veg_fruit", "pkt", r"mushrooms?"),
    ("Cucumber", "veg_fruit", "kg", r"cucumber"),
    ("Tomato", "veg_fruit", "kg", r"tomato"),
    ("Cherry tomato", "veg_fruit", "pkt", r"cherrytomato"),
    ("Corn", "veg_fruit", "pkt", r"corn"),
    ("Vegetables (mixed)", "veg_fruit", "kg", r"vegetables?"),
    ("Rice", "grocery", "kg", r"rice"),
    ("Bread", "grocery", "pkt", r"bread"),
    ("Besan", "grocery", "kg", r"besan"),
    ("Rajma", "grocery", "kg", r"rajma"),
    ("Refined oil", "grocery", "L", r"refined"),
    ("Yeast", "grocery", "pkt", r"yeast"),
    ("Ketchup", "grocery", "btl", r"ketchup"),
    ("Tabasco", "grocery", "btl", r"tabasco|tabisco"),
    ("Peanut butter", "grocery", "pcs", r"peanutbutter"),
    ("Muesli", "grocery", "pkt", r"muesli"),
    ("Maggi", "grocery", "pkt", r"maggi|maggie"),
    ("Fries (frozen)", "grocery", "pkt", r"fries"),
    ("Pepperoni", "grocery", "pkt", r"pepperoni"),
    ("Chicken", "grocery", "kg", r"chicken"),
    ("Straws", "packaging", "pkt", r"straws?"),
    ("Foil", "packaging", "nos", r"foil"),
    ("Wrapping paper", "packaging", "nos", r"wrapping"),
    ("Containers", "packaging", "nos", r"containers?"),
    ("Display packets", "packaging", "nos", r"packets?"),
    ("Tape", "packaging", "nos", r"tape"),
    ("Gloves", "cleaning", "box", r"gloves"),
    ("Kitchen towels", "cleaning", "nos", r"kitchentowel"),
    ("Mop", "cleaning", "nos", r"mop"),
    ("Lizol", "cleaning", "btl", r"lizol"),
    ("Handwash", "cleaning", "btl", r"handwash"),
    ("Cleaning cloth", "cleaning", "nos", r"cloth"),
    ("Brush", "cleaning", "nos", r"brush"),
    ("Fly repellent", "cleaning", "nos", r"makkhi"),
    ("Room freshener", "cleaning", "nos", r"freshener"),
    ("Gas Cylinder", "gas_fuel", "cyl", r"cylinder|gas"),
    ("Batteries", "other", "nos", r"cell|battery|batteries"),
    ("Stationery", "other", "nos", r"markers?|stationery"),
    ("Change given to customer", "other", "nos", r"change"),
]
MIXED = ("Mixed groceries", "grocery", "nos")

# Non-inventory rows: first matching rule on "To + Comment" wins.
OTHER_RULES = [
    (r"^gst\b", ("GST payment", "gst", "nos")),
    (r"\bca\b|notary|filing charges", ("CA & compliance fees", "admin", "nos")),
    (r"anthropic|claude|subs\b", ("Software subscriptions", "admin", "nos")),
    (r"recharge", ("Phone recharge", "utilities", "nos")),
    (r"customer|\bchange\b|extra cash", ("Change given to customer", "other", "nos")),
    (r"\broom\b|\bstay\b", ("Staff room rent", "rent", "nos")),
    (r"\brent\b", ("Rent", "rent", "nos")),
    (r"menu|printing", ("Menu & printing", "marketing", "nos")),
    (r"billboard|sign\b", ("Signage & billboards", "marketing", "nos")),
    (r"team meeting|staff|daaru|alcohol|gift|rashan|whiskey", ("Staff welfare & team meetings", "staff", "nos")),
    (r"agarbatti|dhoopbatti|\bhit\b", ("Agarbatti & pest control", "cleaning", "nos")),
    (r"porter", ("Porter & transport", "transport", "nos")),
    (r"fire|board|pech|hardware", ("Repairs & fittings", "repairs", "nos")),
    (r"electricity", ("Electricity & maintenance", "utilities", "nos")),
]
TYPE_DEFAULTS = {
    "bill": ("Electricity & maintenance", "utilities", "nos"),
    "rent": ("Rent", "rent", "nos"),
    "marketing": ("Signage & billboards", "marketing", "nos"),
}
MISC = ("Miscellaneous", "other", "nos")
# Rows classified as these go to Fixed costs (month = the month paid) instead of daily expenses.
FIXED_HEADS = {
    "Rent": "Rent",
    "Electricity & maintenance": "Electricity bill",
    "Staff room rent": "Staff room rent",
}

OPENING_CATEGORIES = {
    "interior": "interior", "equipments": "equipment", "equipment": "equipment", "misc": "misc",
    "inventory": "inventory", "shop": "shop", "food tasting": "food_tasting", "marketing": "marketing",
}


def text(v):
    if v is None or isinstance(v, datetime):  # a few comment cells were auto-converted to dates
        return ""
    return " ".join(str(v).split())


def classify_inventory(comment):
    s = comment.lower()
    for pattern, token in PHRASES:
        s = re.sub(pattern, f" {token} ", s)
    found = {name for name, _, _, rx in ITEMS if re.search(rf"\b(?:{rx})\b", s)}
    if len(found) == 1:
        name = found.pop()
        return next(i[:3] for i in ITEMS if i[0] == name)
    return MIXED


def classify_other(kind, to, comment):
    s = f"{to} | {comment}".lower()
    for rx, item in OTHER_RULES:
        if re.search(rx, s):
            return item
    return TYPE_DEFAULTS.get(kind, MISC)


def salary_month(paid_on, amount, monthly):
    """Big payments in the first week of a month are last month's salary."""
    if paid_on.day <= 7 and monthly and amount >= monthly / 2:
        return add_months(paid_on.replace(day=1), -1)
    return paid_on.replace(day=1)


# --------------------------------------------------------------- reading a tab

DATE_FORMATS = ["%d-%B-%Y", "%d-%b-%Y", "%d %B %Y", "%d %b %Y", "%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d/%m/%y"]
MODES = [("cash", "cash"), ("upi", "upi"), ("gpay", "upi"), ("phonepe", "upi"), ("paytm", "upi"),
         ("card", "card"), ("bank", "bank"), ("neft", "bank"), ("imps", "bank"), ("transfer", "bank")]


def parse_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = text(v)
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def parse_amount(v):
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return Decimal(str(v)).quantize(Decimal("0.01"))
    s = text(v).replace(",", "").replace("₹", "").strip()
    try:
        return Decimal(s).quantize(Decimal("0.01")) if s else None
    except InvalidOperation:
        return None


def parse_mode(v):
    s = text(v).lower()
    return next((mode for word, mode in MODES if word in s), "")


def read_rows(ws, tab):
    """Rows of one tab as dicts, plus problems ("row 12: ...") for rows that can't be read."""
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], [f"The '{tab}' tab is empty."]
    header = [text(c).lower().rstrip("s") for c in rows[0]]  # "Comments" -> "comment"
    col = {name: header.index(name) for name in ("date", "amount", "by", "to", "type", "comment", "mode") if name in header}
    missing = [n for n in ("date", "amount", "to", "type") if n not in col]
    if missing:
        return [], [f"The '{tab}' tab has no {', '.join(missing)} column (first row: {', '.join(text(c) for c in rows[0] if c)})."]

    def cell(r, name):
        i = col.get(name)
        return r[i] if i is not None and i < len(r) else None

    out, problems = [], []
    for n, r in enumerate(rows[1:], start=2):
        raw_date, raw_amount = cell(r, "date"), cell(r, "amount")
        if text(raw_date) == "" and text(raw_amount) == "" and not isinstance(raw_date, datetime):
            continue
        d, amount = parse_date(raw_date), parse_amount(raw_amount)
        if d is None or amount is None:
            problems.append(f"Row {n}: can't read the date/amount ({text(raw_date) or '—'} / {text(raw_amount) or '—'}).")
            continue
        if not amount:
            continue
        out.append({
            "tab": tab, "row": n, "date": d, "amount": amount,
            "by": text(cell(r, "by")) or BUSINESS_PAYER,
            "to": text(cell(r, "to")),
            "type": text(cell(r, "type")).lower(),
            "comment": text(cell(r, "comment")),
            "mode": parse_mode(cell(r, "mode")),
        })
    return out, problems


# --------------------------------------------------------------- where rows go


def vendor_matcher():
    """to -> Vendor or None, using each vendor's keywords from the Vendors page."""
    vendors = [(v, v.keywords() + [v.name.lower()]) for v in Vendor.objects.filter(is_active=True)]

    def match(to):
        to = to.lower()
        return next((v for v, words in vendors if any(w and w in to for w in words)), None)

    return match


def match_staff(to, monthly_staff):
    """Match the sheet's "To" (often a first or last name) to someone on the Staff page."""
    to = to.lower().strip()
    for s in monthly_staff:
        name = s.name.lower()
        if to and (to == name or to in name.split() or name in to):
            return s
    return None


ADHOC_SALARY = re.compile(r"helper|labou?r|\bdays?\b|daily|ad-?hoc", re.I)
# After opening, equipment and fittings bought for the running café are day-to-day costs.
RUNNING_ITEMS = {
    "interior": ("Repairs & fittings", "repairs", "nos"),
    "equipments": ("Kitchen equipment & tools", "other", "nos"),
    "equipment": ("Kitchen equipment & tools", "other", "nos"),
}


def route(r, vendor_of, monthly_staff, opening_period=False):
    """Decide where a row goes: sets r["dest"] and r["label"] (shown in previews)."""
    kind = r["type"]
    vendor = vendor_of(r["to"]) if kind != "salary" else None
    if vendor:
        r.update(dest="vendor", vendor=vendor.name, label=f"{vendor.name} (vendor)")
    elif kind == "salary" and opening_period:
        staff = match_staff(r["to"], monthly_staff)
        r.update(dest="salary", staff=staff, label=f"Salary · {staff.name if staff else ADHOC[0]}")
    elif kind == "salary" and ADHOC_SALARY.search(f"{r['to']} {r['comment']}"):
        r.update(dest="salary", staff=None, label=f"Salary · {ADHOC[0]}")
    elif kind == "salary":
        r.update(dest="staff_salary", label="Main staff salary (record it on the Staff page)")
    elif kind in RUNNING_ITEMS and not opening_period:
        r.update(dest="expenses", item=RUNNING_ITEMS[kind], label=f"Daily expense · {RUNNING_ITEMS[kind][0]}")
    elif kind in RUNNING_ITEMS:
        category = "interior" if kind == "interior" else "equipment"
        r.update(dest="opening", category=category, label=f"Opening cost · {category.title()}")
    else:
        item = classify_inventory(r["comment"]) if kind == "inventory" else classify_other(kind, r["to"], r["comment"])
        if item[0] in FIXED_HEADS:
            r.update(dest="fixed", head=FIXED_HEADS[item[0]], label=f"Fixed cost · {FIXED_HEADS[item[0]]}")
        else:
            r.update(dest="expenses", item=item, label=f"Daily expense · {item[0]}")
    return r


# --------------------------------------------------------------- saving


class Saver:
    """Creates the rows, tagged with `source` so they can be removed together."""

    def __init__(self, source, user=None):
        self.source, self.user = source, user
        self.payers, self.items, self.heads = {}, {}, {}
        self._adhoc = None

    def payer(self, name):
        key = name.strip().lower()
        if key not in self.payers:
            p = Payer.objects.filter(name__iexact=name.strip()).first()
            if not p:
                p = Payer.objects.create(name=name.strip().title(), is_business=key == BUSINESS_PAYER.lower())
            self.payers[key] = p
        return self.payers[key]

    def item(self, name, category, unit):
        if name not in self.items:
            obj = ExpenseItem.objects.filter(name__iexact=name).first()
            self.items[name] = obj or ExpenseItem.objects.create(name=name, category=category, unit=unit)
        return self.items[name]

    def adhoc(self):
        if self._adhoc is None:
            self._adhoc = Staff.objects.filter(name__iexact=ADHOC[0]).first() or Staff.objects.create(
                name=ADHOC[0], role=ADHOC[1], monthly_salary=0, notes="Paid per day; no fixed monthly salary"
            )
        return self._adhoc

    def save(self, r):
        getattr(self, f"save_{r['dest']}")(r)

    def _common(self, r):
        return {"paid_by": self.payer(r["by"]), "source": self.source, "created_by": self.user}

    def save_expenses(self, r):
        note = ": ".join(x for x in (r["to"], r["comment"]) if x)
        DailyExpense.objects.create(date=r["date"], item=self.item(*r["item"]), amount=r["amount"],
                                    payment_mode=r["mode"], note=note[:200], **self._common(r))

    def save_vendor(self, r):
        vendor = Vendor.objects.get(name=r["vendor"])
        Purchase.objects.create(
            vendor=vendor, bill_date=r["date"], category=vendor.category, total=r["amount"],
            is_paid=True, paid_on=r["date"], payment_mode=r["mode"],
            notes=f"Payment recorded in the expense sheet{': ' + r['comment'] if r['comment'] else ''}"[:300],
            **self._common(r),
        )

    def save_salary(self, r):
        monthly = r["staff"]
        person = monthly or self.adhoc()
        is_advance = "advance" in r["comment"].lower()
        month = r["date"].replace(day=1) if is_advance or not monthly else salary_month(
            r["date"], r["amount"], person.monthly_salary
        )
        SalaryPayment.objects.create(
            staff=person, month=month, kind="advance" if is_advance else "salary", amount=r["amount"],
            paid_on=r["date"], payment_mode=r["mode"],
            note=(": ".join(x for x in (r["to"], r["comment"]) if x) if not monthly else r["comment"])[:200],
            **self._common(r),
        )

    def save_fixed(self, r):
        if r["head"] not in self.heads:
            self.heads[r["head"]] = (FixedCostHead.objects.filter(name__iexact=r["head"]).first()
                                     or FixedCostHead.objects.create(name=r["head"]))
        FixedCost.objects.create(
            head=self.heads[r["head"]], month=r["date"].replace(day=1), amount=r["amount"], paid_on=r["date"],
            payment_mode=r["mode"], note=": ".join(x for x in (r["to"], r["comment"]) if x)[:200], **self._common(r),
        )

    def save_opening(self, r):
        OpeningCost.objects.create(
            date=r["date"], amount=r["amount"], category=r["category"], paid_to=r["to"][:100],
            description=r["comment"][:200], **self._common(r),
        )


# --------------------------------------------------------------- monthly import from the website

# Rows a month import creates; importing the same month again replaces them.
IMPORTED_MODELS = (DailyExpense, FixedCost, SalaryPayment, OpeningCost)
DESTINATIONS = [("expenses", "Daily expenses"), ("fixed", "Fixed costs"), ("salary", "Salaries"), ("opening", "Opening costs")]


class SheetError(Exception):
    pass


def source_for(month):
    return f"gsheet:{month:%Y-%m}"


def sheet_id(url):
    m = re.search(r"/spreadsheets/d/([A-Za-z0-9_-]{20,})", url or "")
    if not m:
        raise SheetError("That doesn't look like a Google Sheets link (it should contain /spreadsheets/d/…).")
    return m.group(1)


def download(url):
    """The whole sheet as .xlsx bytes. The sheet must be shared as 'Anyone with the link can view'."""
    export = f"https://docs.google.com/spreadsheets/d/{sheet_id(url)}/export?format=xlsx"
    req = urllib.request.Request(export, headers={"User-Agent": "Eatalley management"})
    try:
        with urllib.request.urlopen(req, timeout=30) as resp:
            data = resp.read()
    except (urllib.error.URLError, TimeoutError) as e:
        raise SheetError(f"Couldn't download the sheet ({e}). Check the link and try again.") from e
    if not data.startswith(b"PK"):  # Google sends its sign-in page instead of the file
        raise SheetError("Google didn't share the sheet. In Google Sheets, set Share → General access to "
                         "'Anyone with the link' (Viewer).")
    return data


def month_tab(names, month):
    """The tab for this month: 'October', or 'Oct', 'October 2026', 'Oct-26'..."""
    full, short = f"{month:%B}".lower(), f"{month:%b}".lower()
    by_name = {n.strip().lower(): n for n in names}
    if full in by_name:
        return by_name[full]
    for low, name in by_name.items():
        first = re.split(r"[\s\-_/']+", low)[0]
        if first in (full, short) and (str(month.year) in low or str(month.year)[2:] in low or low in (full, short)):
            return name
    raise SheetError(f"The sheet has no '{month:%B}' tab. Its tabs are: {', '.join(names)}.")


def plan_month(url, month):
    """Read this month's tab and decide where every row goes. Nothing is saved."""
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(download(url)), data_only=True, read_only=True)
    tab = month_tab(wb.sheetnames, month)
    rows, problems = read_rows(wb[tab], tab)
    m_end = month_end(month)
    in_month = [r for r in rows if month <= r["date"] <= m_end]
    other_month = [r for r in rows if not month <= r["date"] <= m_end]

    vendor_of = vendor_matcher()
    monthly_staff = list(Staff.objects.filter(monthly_salary__gt=0))
    for r in in_month:
        route(r, vendor_of, monthly_staff)
    to_import = [r for r in in_month if r["dest"] not in ("vendor", "staff_salary")]

    groups = defaultdict(list)
    for r in in_month:
        groups[r["dest"]].append(r)
    digest = hashlib.sha256(
        repr([(r["row"], r["date"], r["amount"], r["by"], r["to"], r["type"], r["comment"], r["mode"]) for r in to_import]).encode()
    ).hexdigest()[:16]
    source = source_for(month)
    return {
        "url": url, "month": month, "tab": tab,
        "rows": to_import, "vendor_rows": groups["vendor"], "other_month": other_month, "problems": problems,
        "summary": [
            {"key": k, "label": label, "n": len(groups[k]), "amount": sum((r["amount"] for r in groups[k]), ZERO)}
            for k, label in DESTINATIONS
        ],
        "total": sum((r["amount"] for r in to_import), ZERO),
        "vendor_total": sum((r["amount"] for r in groups["vendor"]), ZERO),
        "staff_salary_rows": groups["staff_salary"],
        "staff_salary_total": sum((r["amount"] for r in groups["staff_salary"]), ZERO),
        "skipped_n": len(groups["vendor"]) + len(groups["staff_salary"]),
        "digest": digest,
        "previous": SheetImport.objects.filter(month=month).order_by("-created_at").first(),
        "replaces": sum(m.objects.filter(source=source).count() for m in IMPORTED_MODELS),
        # Daily expenses typed into the website for this month: the import leaves them alone.
        "entered_here": DailyExpense.objects.filter(date__range=(month, m_end)).exclude(source=source).count(),
    }


def import_month(plan, user):
    """Save the planned rows, replacing whatever an earlier import of this month created."""
    source = source_for(plan["month"])
    saver = Saver(source, user)
    with transaction.atomic():
        for m in IMPORTED_MODELS:
            m.objects.filter(source=source).delete()
        for r in plan["rows"]:
            saver.save(r)
        log = SheetImport.objects.create(
            month=plan["month"], sheet_url=plan["url"], tab=plan["tab"], rows=len(plan["rows"]),
            total=plan["total"], created_by=user,
        )
    return log
